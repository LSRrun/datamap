#!/usr/bin/env python3
"""Extract the data catalog workbook into a browser-ready JavaScript file."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


HEADERS = {
    "主题域分组": "domain",
    "主题域": "subdomain",
    "业务对象": "object",
    "数据表": "name",
    "英文表名": "englishName",
    "模型负责人": "owner",
    "上线时间": "launchDate",
    "是否线上化": "online",
    "是否入湖": "inLake",
    "说明": "note",
}

REQUIRED_CATALOG_HEADERS = ("主题域分组", "主题域", "业务对象", "数据表", "英文表名")
LEVEL_PREFIX = re.compile(r"^L[1-4]\s*", re.IGNORECASE)
PHYSICAL_SEPARATOR = re.compile(r"[、，,；;|/\r\n]+")
PG_IDENTIFIER = re.compile(r"^[A-Za-z_][A-Za-z0-9_$]*$")


def normalize(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def clean_level(value):
    return LEVEL_PREFIX.sub("", normalize(value)).strip()


def select_catalog_sheet(workbook):
    if "模型目录" in workbook.sheetnames:
        return workbook["模型目录"]
    candidates = []
    for sheet in workbook.worksheets:
        headers = [normalize(cell.value) for cell in sheet[1]]
        if all(name in headers for name in REQUIRED_CATALOG_HEADERS):
            candidates.append(sheet)
    if len(candidates) == 1:
        return candidates[0]
    if not candidates:
        raise ValueError("工作簿中未找到包含必要目录列的工作表")
    raise ValueError("工作簿中存在多个目录工作表，请将目标工作表命名为“模型目录”")


def split_physical_tables(value, default_schema):
    raw_parts = [part.strip() for part in PHYSICAL_SEPARATOR.split(normalize(value)) if part.strip()]
    unique_parts = list(dict.fromkeys(part.upper() for part in raw_parts))
    tables = []
    invalid = []
    defaulted = []

    for raw_name in unique_parts:
        if "." in raw_name:
            pieces = raw_name.split(".", 1)
            schema_name, table_name = pieces[0].strip(), pieces[1].strip()
            schema_defaulted = False
        else:
            schema_name, table_name = default_schema.upper(), raw_name.strip()
            schema_defaulted = True

        if not PG_IDENTIFIER.fullmatch(schema_name) or not PG_IDENTIFIER.fullmatch(table_name):
            invalid.append(raw_name)
            continue

        tables.append({
            "schemaName": schema_name,
            "tableName": table_name,
            "schemaDefaulted": schema_defaulted,
            "rawName": raw_name,
        })
        if schema_defaulted:
            defaulted.append(raw_name)

    return {
        "rawParts": raw_parts,
        "tables": tables,
        "invalid": invalid,
        "defaulted": defaulted,
        "duplicatesRemoved": len(raw_parts) - len(unique_parts),
    }


def extract(workbook_path: Path) -> dict:
    # This workbook's cached worksheet dimension stops at column I even though
    # useful fields continue through column K. Normal mode recalculates the
    # real used range; read-only mode would silently omit "是否入湖" and "说明".
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    sheet = select_catalog_sheet(workbook)
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize(value) for value in next(rows)]
    positions = {name: headers.index(name) for name in HEADERS if name in headers}

    current = {"domain": "", "subdomain": "", "object": ""}
    domain_order: list[str] = []
    tables: list[dict] = []

    for values in rows:
        row = {
            target: normalize(values[positions[source]])
            for source, target in HEADERS.items()
            if source in positions and positions[source] < len(values)
        }
        for level in ("domain", "subdomain", "object"):
            if row.get(level):
                current[level] = row[level]
            else:
                row[level] = current[level]

        if row.get("domain") and row["domain"] not in domain_order:
            domain_order.append(row["domain"])
        if row.get("name"):
            tables.append({key: row.get(key, "") for key in HEADERS.values()})

    counts = {name: 0 for name in domain_order}
    for table in tables:
        counts[table["domain"]] = counts.get(table["domain"], 0) + 1

    return {
        "source": workbook_path.name,
        "domains": [{"name": name, "count": counts.get(name, 0)} for name in domain_order],
        "tables": tables,
    }


def extract_catalog(workbook_path: Path, default_schema: str = "public") -> dict:
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    sheet = select_catalog_sheet(workbook)
    headers = [normalize(cell.value) for cell in sheet[1]]
    missing_headers = [name for name in REQUIRED_CATALOG_HEADERS if name not in headers]
    if missing_headers:
        raise ValueError("工作表缺少必要列：" + "、".join(missing_headers))
    positions = {name: headers.index(name) for name in HEADERS if name in headers}

    current = {"domain": "", "subdomain": "", "object": ""}
    domain_order = []
    assets = []
    last_data_row = 1
    missing_physical = []
    multi_physical = []
    defaulted_schemas = []
    invalid_physical = []
    duplicate_within_row = []
    hierarchy_errors = []

    for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row = {
            target: normalize(values[positions[source]])
            for source, target in HEADERS.items()
            if source in positions and positions[source] < len(values)
        }
        if any(row.values()):
            last_data_row = source_row

        for level in ("domain", "subdomain", "object"):
            if row.get(level):
                current[level] = row[level]
            else:
                row[level] = current[level]

        if row.get("domain"):
            domain_name = clean_level(row["domain"])
            if domain_name and domain_name not in domain_order:
                domain_order.append(domain_name)

        if not row.get("name"):
            continue

        domain = clean_level(row.get("domain"))
        topic = clean_level(row.get("subdomain"))
        business_object = clean_level(row.get("object"))
        name_cn = clean_level(row.get("name"))
        asset_code = f"excel:model-directory:{source_row:04d}"
        missing_levels = [
            label for label, value in (("L1", domain), ("L2", topic), ("L3", business_object), ("L4", name_cn))
            if not value
        ]
        if missing_levels:
            hierarchy_errors.append({"sourceRow": source_row, "missingLevels": missing_levels})

        physical = split_physical_tables(row.get("englishName"), default_schema)
        if not physical["rawParts"]:
            missing_physical.append({"sourceRow": source_row, "assetCode": asset_code, "nameCn": name_cn})
        if len(physical["rawParts"]) > 1:
            multi_physical.append({
                "sourceRow": source_row,
                "assetCode": asset_code,
                "nameCn": name_cn,
                "rawValue": row.get("englishName", ""),
                "parsedCount": len(physical["tables"]),
            })
        if physical["defaulted"]:
            defaulted_schemas.append({
                "sourceRow": source_row,
                "assetCode": asset_code,
                "tables": physical["defaulted"],
                "defaultSchema": default_schema.upper(),
            })
        if physical["invalid"]:
            invalid_physical.append({
                "sourceRow": source_row,
                "assetCode": asset_code,
                "values": physical["invalid"],
            })
        if physical["duplicatesRemoved"]:
            duplicate_within_row.append({
                "sourceRow": source_row,
                "assetCode": asset_code,
                "duplicatesRemoved": physical["duplicatesRemoved"],
            })

        assets.append({
            "assetCode": asset_code,
            "sourceRow": source_row,
            "l1Domain": domain,
            "l2Topic": topic,
            "l3Object": business_object,
            "nameCn": name_cn,
            "owner": normalize(row.get("owner")),
            "description": normalize(row.get("note")),
            "onlineStatus": normalize(row.get("online")),
            "lakeStatus": normalize(row.get("inLake")),
            "launchDate": normalize(row.get("launchDate")),
            "physicalTables": physical["tables"],
        })

    physical_owners = {}
    physical_conflicts = []
    physical_references = 0
    for asset in assets:
        for table in asset["physicalTables"]:
            physical_references += 1
            key = f"{table['schemaName']}.{table['tableName']}"
            owner = physical_owners.get(key)
            if owner and owner["assetCode"] != asset["assetCode"]:
                physical_conflicts.append({
                    "physicalTable": key,
                    "primarySourceRow": owner["sourceRow"],
                    "conflictingSourceRow": asset["sourceRow"],
                    "primaryAssetCode": owner["assetCode"],
                    "conflictingAssetCode": asset["assetCode"],
                })
            else:
                physical_owners[key] = {
                    "assetCode": asset["assetCode"],
                    "sourceRow": asset["sourceRow"],
                }

    domains_with_assets = list(dict.fromkeys(asset["l1Domain"] for asset in assets if asset["l1Domain"]))
    empty_domains = [name for name in domain_order if name not in domains_with_assets]
    workbook.close()

    return {
        "source": {
            "workbook": workbook_path.name,
            "sheet": sheet.title,
            "lastDataRow": last_data_row,
            "defaultSchema": default_schema.upper(),
        },
        "assets": assets,
        "validation": {
            "counts": {
                "scannedRows": max(last_data_row - 1, 0),
                "domainHeadings": len(domain_order),
                "domainsWithAssets": len(domains_with_assets),
                "emptyDomains": len(empty_domains),
                "assets": len(assets),
                "assetsWithoutPhysicalTable": len(missing_physical),
                "rowsWithMultiplePhysicalTables": len(multi_physical),
                "physicalReferences": physical_references,
                "uniquePhysicalTables": len(physical_owners),
                "physicalAssociationConflicts": len(physical_conflicts),
                "rowsUsingDefaultSchema": len(defaulted_schemas),
                "invalidPhysicalRows": len(invalid_physical),
            },
            "emptyDomains": empty_domains,
            "missingPhysicalTables": missing_physical,
            "multiplePhysicalTables": multi_physical,
            "defaultedSchemas": defaulted_schemas,
            "invalidPhysicalTables": invalid_physical,
            "duplicatePhysicalWithinRow": duplicate_within_row,
            "physicalAssociationConflicts": physical_conflicts,
            "hierarchyErrors": hierarchy_errors,
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    parser.add_argument("--format", choices=("browser-js", "catalog-json"), default="browser-js")
    parser.add_argument("--default-schema", default="public")
    args = parser.parse_args()

    if args.format == "catalog-json":
        payload = extract_catalog(args.workbook, args.default_schema)
        print(json.dumps(payload, ensure_ascii=False, separators=(",", ":")))
        return

    if args.output is None:
        parser.error("browser-js 格式需要提供输出文件")
    payload = extract(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        "window.DATA_MAP = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    print(f"Extracted {len(payload['tables'])} tables across {len(payload['domains'])} domains -> {args.output}")


if __name__ == "__main__":
    main()
