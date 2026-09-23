#!/usr/bin/env python3
"""Convert 数据模型全.xlsx into browser-ready catalog data."""

import argparse
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


FIELDS = {
    "主题域分组": "domain",
    "主题域": "subdomain",
    "业务对象": "object",
    "数据表": "name",
    "英文表名": "englishName",
    "模型负责人": "owner",
    "上线时间": "launchDate",
    "是否线上化": "online",
    "是否入湖": "inLake",
    "说明": "note",
}


def text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def extract(source: Path):
    # Normal mode is intentional: this workbook's cached dimension ends at
    # column I, while useful fields continue through column K.
    workbook = load_workbook(source, read_only=False, data_only=True)
    sheet = workbook["模型目录"]
    rows = sheet.iter_rows(values_only=True)
    header = [text(value) for value in next(rows)]
    indexes = {label: header.index(label) for label in FIELDS if label in header}
    current = {"domain": "", "subdomain": "", "object": ""}
    domain_order = []
    tables = []

    for values in rows:
        record = {
            target: text(values[indexes[label]])
            for label, target in FIELDS.items()
            if label in indexes and indexes[label] < len(values)
        }
        for level in ("domain", "subdomain", "object"):
            if record.get(level):
                current[level] = record[level]
            else:
                record[level] = current[level]
        if record.get("domain") and record["domain"] not in domain_order:
            domain_order.append(record["domain"])
        if record.get("name"):
            tables.append({target: record.get(target, "") for target in FIELDS.values()})

    counts = {domain: 0 for domain in domain_order}
    for table in tables:
        counts[table["domain"]] = counts.get(table["domain"], 0) + 1

    return {
        "source": source.name,
        "domains": [{"name": domain, "count": counts[domain]} for domain in domain_order],
        "tables": tables,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    payload = extract(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        "window.DATA_MAP = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    print(f"已导出 {len(payload['tables'])} 张表、{len(payload['domains'])} 个主题域")


if __name__ == "__main__":
    main()
